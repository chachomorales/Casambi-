"""
Tests de humo: que las 17 rutas respondan lo que deben.

No comprueban la lógica de negocio, sino que la app arranca, sirve y guarda sin
reventar. Son la red de seguridad para la migración a web, que toca `app.py` en
una treintena de sitios.
"""

from __future__ import annotations

import json

import pytest

from conftest import RED_ID


# ── Pantallas ─────────────────────────────────────────────────────────────────

def test_index_lista_las_redes(cliente):
    r = cliente.get("/")
    assert r.status_code == 200
    assert "Red de prueba" in r.get_data(as_text=True)


def test_index_sin_cuentas_redirige_a_ajustes(cliente, credenciales):
    credenciales.cuentas = []
    r = cliente.get("/")
    assert r.status_code == 302
    assert "/ajustes" in r.headers["Location"]


def test_index_sin_cargar_muestra_pantalla_de_carga(app_sin_cargar):
    r = app_sin_cargar.test_client().get("/")
    assert r.status_code == 200
    assert "carga/estado" in r.get_data(as_text=True)


def test_carga_estado_devuelve_json(cliente):
    r = cliente.get("/carga/estado")
    assert r.status_code == 200
    assert set(r.get_json()) == {"active", "label", "done", "total", "error"}


def test_ajustes_lista_cuentas(cliente):
    r = cliente.get("/ajustes")
    assert r.status_code == 200
    assert "Cuenta de prueba" in r.get_data(as_text=True)


def test_network_view_pinta_las_nueve_pestanas(cliente):
    r = cliente.get(f"/network/{RED_ID}")
    assert r.status_code == 200
    html = r.get_data(as_text=True)
    for pestana in ("elementos", "conectividad", "luminarias", "pulsadores",
                    "sensores", "grupos", "escenas", "horarios", "planos"):
        assert f'data-tab="{pestana}"' in html


def test_network_view_clasifica_cada_categoria(cliente):
    html = cliente.get(f"/network/{RED_ID}").get_data(as_text=True)
    for nombre in ("Luminaria pasillo", "Sensor recepción",
                   "Pulsador entrada", "Pasarela"):
        assert nombre in html


def test_network_desconocida_da_404(cliente):
    assert cliente.get("/network/no-existe").status_code == 404


def test_logos_se_sirven(cliente):
    assert cliente.get("/logos/impelsa_logo.png").status_code == 200


# ── Cuentas ───────────────────────────────────────────────────────────────────

def test_crear_cuenta(cliente, credenciales):
    r = cliente.post("/ajustes/cuentas", data={
        "label": "Nueva", "api_key": "k", "email": "a@b.c", "password": "x",
    })
    assert r.status_code == 302
    assert any(c["label"] == "Nueva" for c in credenciales.cuentas)


def test_editar_cuenta(cliente, credenciales):
    r = cliente.post("/ajustes/cuentas/cuenta1", data={
        "label": "Renombrada", "api_key": "k2", "email": "a@b.c", "password": "",
    })
    assert r.status_code == 302
    assert credenciales.cuentas[0]["label"] == "Renombrada"


def test_borrar_cuenta(cliente, credenciales):
    r = cliente.post("/ajustes/cuentas/cuenta1", data={"accion": "borrar"})
    assert r.status_code == 302
    assert credenciales.cuentas == []


def test_refresh_descarta_la_cache(cliente):
    import app as app_module
    r = cliente.post(f"/network/{RED_ID}/refresh")
    assert r.status_code == 302
    assert RED_ID not in app_module._state["cache"]


def test_refresh_ya_no_acepta_get(cliente):
    """Como GET, una <img src> remota invalidaba la caché de toda la oficina."""
    assert cliente.get(f"/network/{RED_ID}/refresh").status_code == 405


# ── Anotaciones del usuario ───────────────────────────────────────────────────

def test_guardar_nivel_de_escena(cliente, datos_limpios):
    r = cliente.post(f"/network/{RED_ID}/scene_levels",
                     json={"scene_id": "20", "unit_id": "1", "level": "70"})
    assert r.status_code == 200 and r.get_json()["ok"] is True
    guardado = json.loads((datos_limpios / f"scene_levels_{RED_ID}.json").read_text())
    assert guardado["20"]["1"] == "70"


def test_guardar_botones(cliente, datos_limpios):
    r = cliente.post(f"/network/{RED_ID}/buttons",
                     json={"unit_id": "3", "count": 2})
    assert r.status_code == 200 and r.get_json()["ok"] is True
    guardado = json.loads((datos_limpios / f"buttons_{RED_ID}.json").read_text())
    assert guardado["3"]["count"] == 2


def test_guardar_sensor(cliente, datos_limpios):
    r = cliente.post(f"/network/{RED_ID}/sensors",
                     json={"unit_id": "2", "modo": "Presencia"})
    assert r.status_code == 200 and r.get_json()["ok"] is True
    guardado = json.loads((datos_limpios / f"sensors_{RED_ID}.json").read_text())
    assert guardado["2"]["modo"] == "Presencia"


def test_crear_actualizar_y_borrar_horario(cliente, datos_limpios):
    ruta = datos_limpios / f"schedules_{RED_ID}.json"

    r = cliente.post(f"/network/{RED_ID}/schedules", json={"action": "create"})
    assert r.status_code == 200
    hid = r.get_json()["id"]
    assert len(json.loads(ruta.read_text())) == 1

    r = cliente.post(f"/network/{RED_ID}/schedules", json={
        "action": "update", "id": hid, "nombre": "Apagado nocturno",
        "apagado": "22:00", "escena": "Escena general",
    })
    assert r.status_code == 200
    assert json.loads(ruta.read_text())[0]["nombre"] == "Apagado nocturno"

    r = cliente.post(f"/network/{RED_ID}/schedules",
                     json={"action": "delete", "id": hid})
    assert r.status_code == 200
    assert json.loads(ruta.read_text()) == []


def test_horario_sin_id_da_400(cliente):
    r = cliente.post(f"/network/{RED_ID}/schedules", json={"action": "update"})
    assert r.status_code == 400


# ── Planos ────────────────────────────────────────────────────────────────────

def _png_de_prueba(ancho=40, alto=30) -> bytes:
    import io

    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (ancho, alto), (200, 210, 220)).save(buf, format="PNG")
    return buf.getvalue()


def test_subir_plano_y_servirlo(cliente, datos_limpios):
    import io

    r = cliente.post(
        f"/network/{RED_ID}/planos/upload",
        data={"file": (io.BytesIO(_png_de_prueba()), "planta.png")},
        content_type="multipart/form-data",
    )
    assert r.status_code == 302

    planos = json.loads((datos_limpios / f"planos_{RED_ID}.json").read_text())
    assert len(planos) == 1
    assert planos[0]["name"] == "planta"

    r = cliente.get(f"/network/{RED_ID}/planos/img/{planos[0]['id']}")
    assert r.status_code == 200
    assert r.data[:8] == b"\x89PNG\r\n\x1a\n"


def test_colocar_y_quitar_marcador(cliente, datos_limpios):
    import io

    cliente.post(
        f"/network/{RED_ID}/planos/upload",
        data={"file": (io.BytesIO(_png_de_prueba()), "planta.png")},
        content_type="multipart/form-data",
    )
    pid = json.loads((datos_limpios / f"planos_{RED_ID}.json").read_text())[0]["id"]

    r = cliente.post(f"/network/{RED_ID}/planos", json={
        "action": "marker_set", "plano_id": pid, "unit_id": "1", "x": 0.5, "y": 0.5,
    })
    assert r.status_code == 200 and r.get_json()["ok"] is True
    planos = json.loads((datos_limpios / f"planos_{RED_ID}.json").read_text())
    assert "1" in planos[0]["markers"]

    r = cliente.post(f"/network/{RED_ID}/planos", json={
        "action": "marker_remove", "plano_id": pid, "unit_id": "1",
    })
    assert r.status_code == 200
    planos = json.loads((datos_limpios / f"planos_{RED_ID}.json").read_text())
    assert planos[0]["markers"] == {}


def test_plano_inexistente_da_404(cliente):
    assert cliente.get(f"/network/{RED_ID}/planos/img/999").status_code == 404


# ── Informe ───────────────────────────────────────────────────────────────────

def test_descargar_excel(cliente):
    import io

    from openpyxl import load_workbook

    r = cliente.get(f"/network/{RED_ID}/excel")
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["Content-Type"]

    # Comprobar las 12 hojas y no solo que el zip abre: el informe es el
    # producto final, y una hoja que deje de generarse pasaría inadvertida.
    wb = load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames == [
        "Portada", "Red", "Conectividad", "Elementos", "Luminarias",
        "Pulsadores", "Sensores", "Grupos", "Escenas", "Horarios", "Bitácora",
        "Planos",
    ]
    # La hoja de conectividad debe traer las 4 unidades, no solo cabeceras.
    assert wb["Conectividad"].max_row > 10


# ── Escenas sobre hardware real ───────────────────────────────────────────────

def test_capturar_escena(cliente, app_cargada, datos_limpios, monkeypatch):
    # La ruta duerme 10 s esperando a que las luminarias apliquen la escena.
    monkeypatch.setattr("app.time.sleep", lambda _s: None)

    r = cliente.post(f"/network/{RED_ID}/scenes/20/capture")
    assert r.status_code == 200
    cuerpo = r.get_json()
    assert cuerpo["ok"] is True
    assert cuerpo["captured"]["1"] == "80"
    assert app_cargada.cliente_api_falso.escenas_activadas == [(RED_ID, "20")]

    guardado = json.loads((datos_limpios / f"scene_levels_{RED_ID}.json").read_text())
    assert guardado["20"]["1"] == "80"


def test_capturar_escena_inexistente(cliente, monkeypatch):
    monkeypatch.setattr("app.time.sleep", lambda _s: None)
    r = cliente.post(f"/network/{RED_ID}/scenes/999/capture")
    assert r.status_code == 404

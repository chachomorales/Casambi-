"""
Temperatura de las luminarias que la reportan (hoy, las Eulum TRED-E-CSB-2A).

Llega en la lista `sensors` del estado, no en `controls`, con la hora de la
última lectura. Lo que se comprueba: que se lee, que se enseña siempre con su
fecha, y que las redes sin ella no ganan una columna vacía.
"""

from __future__ import annotations

import io
from datetime import datetime, timedelta

import pytest
from openpyxl import load_workbook

import app as app_module
import config
import report
from conftest import RED_ID, STATE


def _ms(momento: datetime) -> int:
    """Hora local de la app → epoch en milisegundos, como el `timestamp` de la API."""
    return int(momento.replace(tzinfo=config.zona_horaria()).timestamp() * 1000)


@pytest.fixture
def con_temperatura(app_cargada):
    """La luminaria 1 reporta 35 °C, leídos hace dos horas."""
    leida = config.ahora().replace(second=0, microsecond=0) - timedelta(hours=2)
    estado = {**STATE, "units": [dict(u) for u in STATE["units"]]}
    estado["units"][0]["sensors"] = [
        {"name": "Temperature", "value": 35.0, "timestamp": _ms(leida)}]
    app_module._state["cache"][RED_ID]["state"] = estado
    return leida


def test_lee_la_temperatura_con_su_hora():
    leida = datetime(2026, 9, 27, 14, 3)
    t = report.temperatura_unidad(
        {"sensors": [{"name": "Temperature", "value": 35.0, "timestamp": _ms(leida)}]})
    assert t == {"grados": 35, "leida": leida}
    assert report.texto_temperatura(t) == "35 °C · 27/09/2026 14:03"


def test_sin_sensores_no_hay_temperatura():
    assert report.temperatura_unidad({"controls": [[{"type": "Dimmer", "value": 1}]]}) is None
    assert report.temperatura_unidad({"sensors": [{"name": "Lux", "value": 300}]}) is None
    assert report.texto_temperatura(None) == "-"


def test_hace_en_palabras(monkeypatch):
    ahora = datetime(2026, 9, 27, 12, 0)
    monkeypatch.setattr(config, "ahora", lambda: ahora)
    assert app_module.hace(ahora + timedelta(seconds=30)) == "ahora"   # reloj adelantado
    assert app_module.hace(ahora - timedelta(minutes=12)) == "hace 12 min"
    assert app_module.hace(ahora - timedelta(hours=10)) == "hace 10 h"
    assert app_module.hace(ahora - timedelta(days=3)) == "hace 3 días"


def test_pestanas_muestran_la_temperatura(cliente, con_temperatura):
    html = cliente.get(f"/network/{RED_ID}").get_data(as_text=True)
    assert html.count("<th>Temperatura</th>") == 2        # Conectividad y Luminarias
    assert "35 °C" in html
    assert "hace 2 h" in html
    assert f"{con_temperatura:%d/%m/%Y %H:%M}" in html      # la fecha exacta, en el title


def test_excel_lleva_la_temperatura(cliente, con_temperatura):
    wb = load_workbook(io.BytesIO(cliente.get(f"/network/{RED_ID}/excel").data))
    esperado = f"35 °C · {con_temperatura:%d/%m/%Y %H:%M}"

    lum = wb["Luminarias"]
    assert lum.cell(row=1, column=8).value == "Temperatura"
    assert lum.cell(row=2, column=8).value == esperado

    con = wb["Conectividad"]
    valores = [c.value for fila in con.iter_rows() for c in fila]
    assert "Temperatura" in valores and esperado in valores
    assert con.auto_filter.ref.startswith("A17:L")


def test_red_sin_temperatura_no_gana_columnas(cliente):
    html = cliente.get(f"/network/{RED_ID}").get_data(as_text=True)
    assert "<th>Temperatura</th>" not in html
    wb = load_workbook(io.BytesIO(cliente.get(f"/network/{RED_ID}/excel").data))
    assert "Temperatura" not in [c.value for c in wb["Luminarias"][1]]
    assert wb["Conectividad"].auto_filter.ref.startswith("A17:K")

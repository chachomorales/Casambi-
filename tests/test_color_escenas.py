"""
Color de las luminarias RGBW en las escenas.

La API no lo guarda en la escena: se captura activándola, igual que la
intensidad. Lo delicado es el cambio de perfil: una CBU-PWM4 puede pasar de
RGBW a un canal regulable sin cambiar de ID, y el color capturado antes no
debe seguir apareciendo como si fuera verdad.
"""

from __future__ import annotations

import io
import json

import pytest
from openpyxl import load_workbook

import app as app_module
from conftest import FIXTURES, RED_ID

# Tal cual lo devuelve la API, con los espacios irregulares dentro de rgb()
ESTADO_RGBW = {
    "gateway": {"name": "Pasarela"},
    "units": [
        {"id": 1, "name": "Luminaria pasillo", "online": True, "status": "ok",
         "dimLevel": 0.8, "activeSceneId": 20,
         "controls": [[
             {"type": "Dimmer", "value": 0.8},
             {"type": "Color", "hue": 0.788, "sat": 1.0, "rgb": "rgb(183,  0, 255)"},
             {"type": "White", "value": 0.5},
         ]]},
        {"id": 2, "name": "Sensor recepción", "online": True, "status": "ok",
         "dimLevel": 0.5, "activeSceneId": 20},
    ],
}

PERFIL_RGBW = {"vendor": "Casambi", "model": "CBU-PWM4 RGBW", "controls": [
    {"type": "dimmer", "id": 0}, {"type": "rgb", "id": 8}, {"type": "white", "id": 26},
]}
PERFIL_DIMMER = {"vendor": "Casambi", "model": "CBU-PWM4 1ch", "controls": [
    {"type": "dimmer", "id": 0},
]}


@pytest.fixture
def perfil(app_cargada):
    """Cambia el perfil de la unidad 1 (fixture 100) en la caché de la red."""
    def poner(fixture):
        fixtures = dict(FIXTURES)
        fixtures[100] = fixture
        app_module._state["cache"][RED_ID]["fixtures"] = fixtures
    return poner


@pytest.fixture
def capturar(cliente, app_cargada, monkeypatch):
    monkeypatch.setattr("app.time.sleep", lambda _s: None)

    def hacer(estado):
        monkeypatch.setattr(app_cargada.cliente_api_falso, "get_network_state",
                            lambda _nid: estado)
        r = cliente.post(f"/network/{RED_ID}/scenes/20/capture")
        assert r.status_code == 200
        return r.get_json()
    return hacer


def _hoja_escenas(cliente):
    r = cliente.get(f"/network/{RED_ID}/excel")
    assert r.status_code == 200
    return load_workbook(io.BytesIO(r.data))["Escenas"]


def test_lee_color_y_blanco_del_estado():
    assert app_module._color_de_estado(ESTADO_RGBW["units"][0]) == {
        "rgb": "#B700FF", "blanco": 50}
    assert app_module._color_de_estado(ESTADO_RGBW["units"][1]) is None


def test_captura_guarda_el_color(capturar, perfil, cliente, datos_limpios):
    perfil(PERFIL_RGBW)
    cuerpo = capturar(ESTADO_RGBW)
    assert cuerpo["colores"] == {"1": {"rgb": "#B700FF", "blanco": 50}}

    guardado = json.loads((datos_limpios / f"scene_colors_{RED_ID}.json").read_text())
    assert guardado == {"20": {"1": {"rgb": "#B700FF", "blanco": 50}}}

    html = cliente.get(f"/network/{RED_ID}").get_data(as_text=True)
    assert "background: #B700FF" in html
    assert "Blanco 50 %" in html

    ws = _hoja_escenas(cliente)
    cabecera = [c.value for c in ws[1]]
    assert cabecera[6:8] == ["Color", "Blanco"]
    fila = next(r for r in ws.iter_rows(min_row=2) if r[4].value == "Luminaria pasillo")
    assert fila[6].value == "#B700FF"
    assert fila[6].fill.fgColor.rgb.endswith("B700FF")
    assert fila[7].value == "50 %"


def test_columna_color_sale_antes_de_capturar(perfil, cliente):
    """El perfil basta para saber que hay color; sin él no hay columna."""
    perfil(PERFIL_RGBW)
    assert "<th>Color</th>" in cliente.get(f"/network/{RED_ID}").get_data(as_text=True)
    perfil(PERFIL_DIMMER)
    assert "<th>Color</th>" not in cliente.get(f"/network/{RED_ID}").get_data(as_text=True)


def test_cambio_de_perfil_oculta_el_color_capturado(capturar, perfil, cliente):
    perfil(PERFIL_RGBW)
    capturar(ESTADO_RGBW)

    # En la app de Casambi pasan la luminaria a un canal regulable
    perfil(PERFIL_DIMMER)
    html = cliente.get(f"/network/{RED_ID}").get_data(as_text=True)
    assert "#B700FF" not in html
    ws = _hoja_escenas(cliente)
    assert "Color" not in [c.value for c in ws[1]]
    assert "#B700FF" not in [c.value for fila in ws.iter_rows() for c in fila]


def test_recaptura_sin_color_lo_borra(capturar, perfil, datos_limpios):
    perfil(PERFIL_RGBW)
    capturar(ESTADO_RGBW)

    sin_color = json.loads(json.dumps(ESTADO_RGBW))
    sin_color["units"][0]["controls"] = [[{"type": "Dimmer", "value": 0.8}]]
    cuerpo = capturar(sin_color)
    assert cuerpo["colores"] == {}
    guardado = json.loads((datos_limpios / f"scene_colors_{RED_ID}.json").read_text())
    assert guardado == {}


def test_red_sin_color_no_cambia_el_excel(cliente):
    ws = _hoja_escenas(cliente)
    assert [c.value for c in ws[1]] == [
        "ID", "Nombre", "Tipo", "Nº Dispositivos", "Dispositivo", "Intensidad"]

"""
Medidas de los módulos de medición de energía (hoy, el McWong
PSC-WCM-450-Power-Metering, fixture 31636).

Llegan en la lista `sensors` del estado, como la temperatura, con la hora de la
última lectura. Lo que se comprueba: que se leen con sus unidades, que nunca se
enseñan sin fecha, que el Excel las da como números, y que las redes sin ellas
no ganan columnas vacías.
"""

from __future__ import annotations

import io
from datetime import datetime, timedelta

import pytest
from openpyxl import load_workbook

import app as app_module
import config
import report
from conftest import RED_ID, STATE


def _ms(momento: datetime) -> int:
    """Hora local de la app → epoch en milisegundos, como el `timestamp` de la API."""
    return int(momento.replace(tzinfo=config.zona_horaria()).timestamp() * 1000)


def _sensores(leida: datetime, **valores) -> list:
    return [{"name": n, "value": v, "timestamp": _ms(leida)} for n, v in valores.items()]


@pytest.fixture
def con_medida(app_cargada):
    """La luminaria 1 mide 230 V, 180 mA, 41 W y 1250 Wh, leídos hace 5 min."""
    leida = config.ahora().replace(second=0, microsecond=0) - timedelta(minutes=5)
    estado = {**STATE, "units": [dict(u) for u in STATE["units"]]}
    estado["units"][0]["sensors"] = _sensores(
        leida, Energy=1250.0, Voltage=230.0, Current=180.0, Power=41.0)
    app_module._state["cache"][RED_ID]["state"] = estado
    return leida


def test_lee_las_cuatro_medidas_con_su_hora():
    leida = datetime(2026, 10, 1, 16, 52)
    m = report.medida_electrica({"sensors": _sensores(
        leida, Energy=0.0, Voltage=0.0, Current=0.0, Power=0.0)})
    assert m == {"tension": 0, "corriente": 0, "potencia": 0, "energia": 0, "leida": leida}
    assert report.texto_medida(m) == "0 V · 0 mA · 0 W · 0 Wh · 01/10/2026 16:52"
    assert report.texto_medida(m, con_fecha=False) == "0 V · 0 mA · 0 W · 0 Wh"


def test_con_horas_distintas_vale_la_mas_vieja():
    vieja, nueva = datetime(2026, 10, 1, 9, 0), datetime(2026, 10, 1, 16, 0)
    m = report.medida_electrica({"sensors": _sensores(vieja, Power=40.0)
                                 + _sensores(nueva, Energy=900.0)})
    assert m["leida"] == vieja


def test_sin_medidas_no_hay_nada():
    assert report.medida_electrica({"controls": [[{"type": "Dimmer", "value": 1}]]}) is None
    # La temperatura de las Eulum va en la misma lista y no es una medida eléctrica
    assert report.medida_electrica({"sensors": [{"name": "Temperature", "value": 35}]}) is None
    assert report.texto_medida(None) == "-"


def test_pestana_luminarias_muestra_la_medida(cliente, con_medida):
    html = cliente.get(f"/network/{RED_ID}").get_data(as_text=True)
    assert html.count("<th>Medición eléctrica</th>") == 1
    assert "230 V · 180 mA · 41 W · 1250 Wh" in html
    assert "hace 5 min" in html
    assert f"{con_medida:%d/%m/%Y %H:%M}" in html      # la fecha exacta, en el title


def test_excel_lleva_la_medida_en_numeros(cliente, con_medida):
    wb = load_workbook(io.BytesIO(cliente.get(f"/network/{RED_ID}/excel").data))
    lum = wb["Luminarias"]
    cabecera = [c.value for c in lum[1]]
    i = cabecera.index("Tensión (V)") + 1
    assert cabecera[i - 1:i + 4] == ["Tensión (V)", "Corriente (mA)", "Potencia (W)",
                                     "Energía (Wh)", "Lectura eléctrica"]
    fila = [lum.cell(row=2, column=i + k).value for k in range(5)]
    assert fila == [230, 180, 41, 1250, f"{con_medida:%d/%m/%Y %H:%M}"]


def test_red_sin_medidas_no_gana_columnas(cliente):
    html = cliente.get(f"/network/{RED_ID}").get_data(as_text=True)
    assert "Medición eléctrica" not in html
    wb = load_workbook(io.BytesIO(cliente.get(f"/network/{RED_ID}/excel").data))
    assert "Tensión (V)" not in [c.value for c in wb["Luminarias"][1]]

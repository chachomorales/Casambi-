"""
Presencia y luz que miden los sensores (p. ej. los STARCO PIR + fotosensor).

Llegan en `controls` del estado, sin hora propia. Lo que se comprueba: que se
leen, que un sensor offline no enseña su último valor como si fuera actual, y
que las redes sin ellos no ganan columnas vacías, ni en la pestaña ni en el
Excel.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

import app as app_module
import report
from conftest import RED_ID, STATE


def _vivo(*controles, online=True) -> dict:
    return {"online": online, "controls": [list(controles)]}


PRESENTE = {"name": "presence", "type": "Presence", "status": "present"}
AUSENTE = {"name": "presence", "type": "Presence", "status": "absent"}
LUX = {"name": "lux", "type": "Lux", "value": 37.4}


def test_lee_presencia_y_luz():
    assert report.lectura_sensor(_vivo(PRESENTE, LUX)) == {"presencia": True, "lux": 37}
    assert report.lectura_sensor(_vivo(AUSENTE)) == {"presencia": False, "lux": None}


def test_sensor_offline_no_da_lectura():
    """Conserva el último valor; enseñarlo lo haría pasar por actual."""
    assert report.lectura_sensor(_vivo(AUSENTE, LUX, online=False)) is None


def test_sin_controles_de_sensor_no_hay_lectura():
    assert report.lectura_sensor(_vivo({"type": "Dimmer", "value": 1})) is None
    assert report.lectura_sensor({"online": True, "controls": [[]]}) is None
    assert report.lectura_sensor({}) is None


@pytest.fixture
def con_lectura(app_cargada):
    """El sensor 2 detecta presencia y mide 37 lx."""
    estado = {**STATE, "units": [dict(u) for u in STATE["units"]]}
    estado["units"][1]["controls"] = [[PRESENTE, LUX]]
    app_module._state["cache"][RED_ID]["state"] = estado


def test_pestana_sensores_muestra_la_lectura(cliente, con_lectura):
    html = cliente.get(f"/network/{RED_ID}").get_data(as_text=True)
    assert "<th>Presencia</th>" in html and "<th>Luz</th>" in html
    assert ">Presente</td>" in html
    assert ">37 lx</td>" in html
    assert "Lectura al descargar la red" in html


def test_red_sin_lecturas_no_gana_columnas(cliente):
    html = cliente.get(f"/network/{RED_ID}").get_data(as_text=True)
    assert "<th>Presencia</th>" not in html
    assert "<th>Luz</th>" not in html


def test_excel_lleva_la_lectura(cliente, con_lectura):
    wb = load_workbook(io.BytesIO(cliente.get(f"/network/{RED_ID}/excel").data))
    ws = wb["Sensores"]
    cabecera = [c.value for c in ws[1]]
    assert cabecera[-2:] == ["Presencia", "Luz (lx)"]
    fila = [c.value for c in ws[2]]
    assert fila[-2:] == ["Presente", 37]
    notas = [c.value for c in ws["A"] if isinstance(c.value, str)]
    assert any(n.startswith("Presencia y luz: lectura al descargar la red") for n in notas)
    assert ws.column_dimensions["A"].width < 20          # la nota no ensancha el ID


def test_excel_sin_lecturas_no_gana_columnas(cliente):
    wb = load_workbook(io.BytesIO(cliente.get(f"/network/{RED_ID}/excel").data))
    cabecera = [c.value for c in wb["Sensores"][1]]
    assert "Presencia" not in cabecera and "Luz (lx)" not in cabecera
